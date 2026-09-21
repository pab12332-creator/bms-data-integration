from docx import Document

def llenar_tablas_reporte(ruta_docx, datos):
    """
    Abre el documento Word generado por docxtpl y llena las tablas 
    dinámicas usando python-docx para asegurar el crecimiento vertical.
    """
    doc = Document(ruta_docx)
    
    # IMPORTANTE: Los índices de las tablas en python-docx empiezan en 0.
    # Asumiendo tu estructura:
    # doc.tables[0] = Datos Generales
    # doc.tables[1] = Inspección
    # doc.tables[2] = Parámetros Eléctricos
    # doc.tables[3] = Matriz de Dispositivos
    # doc.tables[4] = Detalle de Fallas

    # --- TABLA 3: Parámetros Eléctricos (Índice 2) ---
    if len(doc.tables) > 2:
        tabla_electricos = doc.tables[2]
        for tablero in datos.tableros_electricos:
            row_cells = tabla_electricos.add_row().cells
            row_cells[0].text = tablero.id
            row_cells[1].text = tablero.voltaje_ln
            row_cells[2].text = tablero.voltaje_ll
            row_cells[3].text = tablero.amperaje
            row_cells[4].text = tablero.estatus
            
    # --- TABLA 4: Matriz de Dispositivos (Índice 3) ---
    if len(doc.tables) > 3:
        tabla_dispositivos = doc.tables[3]
        for item in datos.matriz_dispositivos:
            row_cells = tabla_dispositivos.add_row().cells
            row_cells[0].text = item.ubicacion_tablero
            row_cells[1].text = item.tipo_dispositivo
            row_cells[2].text = str(item.cantidad)
            row_cells[3].text = item.voltaje
            row_cells[4].text = item.temperatura
            row_cells[5].text = item.comms_str
            row_cells[6].text = item.torque_str
            row_cells[7].text = item.estatus_str
            
    # --- TABLA 5: Detalle de Fallas (Índice 4) ---
    if len(doc.tables) > 4:
        tabla_fallas = doc.tables[4]
        if datos.lista_fallas and len(datos.lista_fallas) > 0:
            for falla in datos.lista_fallas:
                row_cells = tabla_fallas.add_row().cells
                row_cells[0].text = falla.componente_id
                row_cells[1].text = falla.descripcion
                row_cells[2].text = falla.accion
        else:
            # Fila por defecto si no hay incidencias (REQUERIMIENTO: Leyenda específica)
            row_cells = tabla_fallas.add_row().cells
            row_cells[0].text = "-"
            row_cells[1].text = "No se reportaron anomalías durante el mantenimiento"
            row_cells[2].text = "-"
            
    # Guardamos los cambios sobre el mismo archivo
    doc.save(ruta_docx)